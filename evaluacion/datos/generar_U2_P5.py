"""
Genera Excel de verificación para Evaluación U2 - Pregunta 5
Tipo de Cambio Real Efectivo (REER) Argentina vs Brasil — 2024
Fuente: FRED / BIS (Bank for International Settlements)
Series: RBARBIS (Argentina), RBBRBIS (Brasil)
"""

import sys
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
import os

# --- DATOS ---
# Descargados de FRED: https://fred.stlouisfed.org/graph/fredgraph.csv?id=RBARBIS,RBBRBIS
# Índice base 2020 = 100, mensual
# Un valor más alto = moneda más apreciada = bienes más caros en USD = menos competitivo

reer_2024 = [
    ('2024-01', 94.77, 119.80),
    ('2024-02', 106.79, 118.34),
    ('2024-03', 116.82, 116.83),
    ('2024-04', 126.15, 114.02),
    ('2024-05', 128.47, 113.48),
    ('2024-06', 133.70, 108.90),
    ('2024-07', 136.52, 105.54),
    ('2024-08', 138.10, 103.79),
    ('2024-09', 139.29, 103.65),
    ('2024-10', 141.37, 103.73),
    ('2024-11', 145.50, 102.35),
    ('2024-12', 149.37, 98.37),
]

# --- CREAR WORKBOOK ---
wb = openpyxl.Workbook()

# Estilos
header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
data_font = Font(name='Calibri', size=11)
title_font = Font(name='Calibri', bold=True, size=14, color='1F4E79')
subtitle_font = Font(name='Calibri', bold=True, size=11, color='E8833A')
result_fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
result_font = Font(name='Calibri', bold=True, size=12, color='27AE60')
arg_fill = PatternFill(start_color='E8D4F0', end_color='E8D4F0', fill_type='solid')
bra_fill = PatternFill(start_color='D4E8F0', end_color='D4E8F0', fill_type='solid')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

def style_data(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = data_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

# ============================================================
# HOJA 1: DATOS
# ============================================================
ws1 = wb.active
ws1.title = 'Datos'

ws1.cell(row=1, column=1, value='Evaluación U2 — Pregunta 5: Tipo de Cambio Real Efectivo (REER)').font = title_font
ws1.cell(row=2, column=1, value='Fuente: FRED / BIS — Índice base 2020=100, mensual').font = subtitle_font
ws1.cell(row=3, column=1, value='Interpretación: valor más ALTO = moneda más apreciada = bienes más caros en USD = MENOS competitivo').font = Font(name='Calibri', size=10, italic=True, color='E74C3C')

# Tabla mensual 2024
row = 5
headers = ['Mes', 'Argentina (RBARBIS)', 'Brasil (RBBRBIS)', 'Brecha (ARG − BRA)']
for i, h in enumerate(headers):
    ws1.cell(row=row, column=i+1, value=h)
style_header(ws1, row, len(headers))

for i, (mes, arg, bra) in enumerate(reer_2024):
    r = row + 1 + i
    ws1.cell(row=r, column=1, value=mes)
    ws1.cell(row=r, column=2, value=arg)
    ws1.cell(row=r, column=3, value=bra)
    ws1.cell(row=r, column=4, value=round(arg - bra, 1))
    style_data(ws1, r, len(headers))
    ws1.cell(row=r, column=2).number_format = '0.0'
    ws1.cell(row=r, column=3).number_format = '0.0'
    ws1.cell(row=r, column=4).number_format = '0.0'
    # Color condicional en brecha
    brecha = arg - bra
    if brecha > 0:
        ws1.cell(row=r, column=4).font = Font(name='Calibri', size=11, color='E74C3C', bold=True)
    else:
        ws1.cell(row=r, column=4).font = Font(name='Calibri', size=11, color='27AE60', bold=True)

ws1.column_dimensions['A'].width = 12
ws1.column_dimensions['B'].width = 22
ws1.column_dimensions['C'].width = 22
ws1.column_dimensions['D'].width = 22

# URLs de descarga
row_url = 19
ws1.cell(row=row_url, column=1, value='URLs para estudiantes:').font = subtitle_font
ws1.cell(row=row_url+1, column=1, value='Argentina REER:')
ws1.cell(row=row_url+1, column=2, value='https://fred.stlouisfed.org/series/RBARBIS')
ws1.cell(row=row_url+2, column=1, value='Brasil REER:')
ws1.cell(row=row_url+2, column=2, value='https://fred.stlouisfed.org/series/RBBRBIS')
ws1.cell(row=row_url+3, column=1, value='Ambos en CSV:')
ws1.cell(row=row_url+3, column=2, value='https://fred.stlouisfed.org/graph/fredgraph.csv?id=RBARBIS,RBBRBIS&cosd=2024-01-01&coed=2024-12-01')

# ============================================================
# HOJA 2: CÁLCULOS
# ============================================================
ws2 = wb.create_sheet('Cálculos')

ws2.cell(row=1, column=1, value='Cálculos de verificación').font = title_font

# Promedios
arg_values = [r[1] for r in reer_2024]
bra_values = [r[2] for r in reer_2024]
avg_arg = sum(arg_values) / len(arg_values)
avg_bra = sum(bra_values) / len(bra_values)
brecha_avg = avg_arg - avg_bra

ws2.cell(row=3, column=1, value='Cálculo').font = Font(bold=True)
ws2.cell(row=3, column=2, value='Valor').font = Font(bold=True)
ws2.cell(row=3, column=3, value='Detalle').font = Font(bold=True)
style_header(ws2, 3, 3)

calculos = [
    ('Promedio Argentina 2024', round(avg_arg, 1), f'Promedio de 12 valores mensuales'),
    ('Promedio Brasil 2024', round(avg_bra, 1), f'Promedio de 12 valores mensuales'),
    ('Brecha promedio (ARG − BRA)', round(brecha_avg, 1), f'{round(avg_arg, 1)} − {round(avg_bra, 1)}'),
    ('Mínimo Argentina (ene-24)', min(arg_values), 'Post-devaluación dic-2023'),
    ('Máximo Argentina (dic-24)', max(arg_values), 'Máxima apreciación del año'),
    ('Apreciación ARG en el año', f'+{round((max(arg_values) - min(arg_values)) / min(arg_values) * 100, 0)}%', f'De {min(arg_values)} a {max(arg_values)}'),
    ('Brecha mínima (ene-24)', round(reer_2024[0][1] - reer_2024[0][2], 1), 'ARG más barata que BRA'),
    ('Brecha máxima (dic-24)', round(reer_2024[-1][1] - reer_2024[-1][2], 1), 'ARG mucho más cara que BRA'),
]

for i, (label, valor, detalle) in enumerate(calculos):
    r = 4 + i
    ws2.cell(row=r, column=1, value=label)
    ws2.cell(row=r, column=2, value=valor)
    ws2.cell(row=r, column=3, value=detalle)
    style_data(ws2, r, 3)

# Resultado destacado
r_result = 13
ws2.cell(row=r_result, column=1, value='RESULTADO:').font = result_font
ws2.cell(row=r_result, column=2, value=f'Brecha promedio = +{round(brecha_avg, 1)} puntos').font = result_font
for c in range(1, 4):
    ws2.cell(row=r_result, column=c).fill = result_fill

ws2.cell(row=15, column=1, value='Interpretación:').font = subtitle_font
ws2.cell(row=16, column=1, value='Argentina promedió ~130 y Brasil ~109 en 2024.')
ws2.cell(row=17, column=1, value='Argentina se apreció 58% en el año (de 94,8 a 149,4).')
ws2.cell(row=18, column=1, value='Brasil se depreció 18% (de 119,8 a 98,4).')
ws2.cell(row=19, column=1, value='La brecha pasó de -25 (ARG barata) a +51 (ARG 51 puntos más cara).')
ws2.cell(row=20, column=1, value='Esto explica por qué las pechugas brasileñas llegan 30% más baratas.')

ws2.cell(row=22, column=1, value='RESPUESTA CORRECTA: a)').font = Font(bold=True, size=12, color='1F4E79')
ws2.cell(row=23, column=1, value='Argentina ~130, Brasil ~109, brecha ~21 puntos.')
ws2.cell(row=24, column=1, value='La apreciación real encarece productos argentinos en dólares.')
ws2.cell(row=25, column=1, value='Conexión con H-O: las dotaciones no cambiaron, pero el TCR distorsionó precios relativos.')

ws2.column_dimensions['A'].width = 35
ws2.column_dimensions['B'].width = 25
ws2.column_dimensions['C'].width = 40

# ============================================================
# HOJA 3: GRÁFICO
# ============================================================
ws3 = wb.create_sheet('Gráfico')

ws3.cell(row=1, column=1, value='REER Argentina vs Brasil — 2024').font = title_font

# Datos para el gráfico
row = 3
headers_g = ['Mes', 'Argentina', 'Brasil']
for i, h in enumerate(headers_g):
    ws3.cell(row=row, column=i+1, value=h)
style_header(ws3, row, 3)

meses_label = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
for i, (fecha, arg, bra) in enumerate(reer_2024):
    r = 4 + i
    ws3.cell(row=r, column=1, value=meses_label[i])
    ws3.cell(row=r, column=2, value=arg)
    ws3.cell(row=r, column=3, value=bra)
    style_data(ws3, r, 3)

# Línea base 100
for i in range(12):
    ws3.cell(row=4+i, column=4, value=100)
ws3.cell(row=3, column=4, value='Base 2020')
ws3.cell(row=3, column=4).font = header_font
ws3.cell(row=3, column=4).fill = PatternFill(start_color='999999', end_color='999999', fill_type='solid')

# Crear gráfico de líneas
chart = LineChart()
chart.title = "Tipo de Cambio Real Efectivo — 2024 (índice 2020=100)"
chart.y_axis.title = "Índice REER (más alto = más caro en USD)"
chart.x_axis.title = "Mes"
chart.style = 10
chart.width = 20
chart.height = 14
chart.y_axis.scaling.min = 85
chart.y_axis.scaling.max = 160

# Series
data_arg = Reference(ws3, min_col=2, min_row=3, max_row=15)
data_bra = Reference(ws3, min_col=3, min_row=3, max_row=15)
data_base = Reference(ws3, min_col=4, min_row=3, max_row=15)
cats = Reference(ws3, min_col=1, min_row=4, max_row=15)

chart.add_data(data_arg, titles_from_data=True)
chart.add_data(data_bra, titles_from_data=True)
chart.add_data(data_base, titles_from_data=True)
chart.set_categories(cats)

# Colores
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.line import LineProperties, LineEndProperties
chart.series[0].graphicalProperties.line.solidFill = "E74C3C"  # Argentina = rojo
chart.series[0].graphicalProperties.line.width = 28000
chart.series[1].graphicalProperties.line.solidFill = "27AE60"  # Brasil = verde
chart.series[1].graphicalProperties.line.width = 28000
chart.series[2].graphicalProperties.line.solidFill = "999999"  # Base = gris
chart.series[2].graphicalProperties.line.dashStyle = "dash"
chart.series[2].graphicalProperties.line.width = 14000

ws3.add_chart(chart, "A17")

# Anotaciones
ws3.cell(row=34, column=1, value='Lectura del gráfico:').font = subtitle_font
ws3.cell(row=35, column=1, value='• Línea ROJA (Argentina): sube de 94,8 a 149,4 → apreciación real masiva (+58%)')
ws3.cell(row=36, column=1, value='• Línea VERDE (Brasil): baja de 119,8 a 98,4 → depreciación real (-18%)')
ws3.cell(row=37, column=1, value='• En enero ARG era más barata; desde abril se invierte; en diciembre la brecha es de +51 puntos')
ws3.cell(row=38, column=1, value='• La zona entre ambas líneas (desde abril) representa la pérdida de competitividad argentina')

ws3.column_dimensions['A'].width = 12
ws3.column_dimensions['B'].width = 12
ws3.column_dimensions['C'].width = 12
ws3.column_dimensions['D'].width = 12

# --- GUARDAR ---
output_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(output_dir, 'U2_P5_tipo_cambio_real.xlsx')
wb.save(output_path)
print(f'Excel generado: {output_path}')
print(f'3 hojas: Datos, Cálculos, Gráfico')
