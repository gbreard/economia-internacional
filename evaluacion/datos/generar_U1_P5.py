"""
Genera U1_P5_cuenta_corriente.xlsx
Datos: Current account balance (% of GDP) para Argentina y Corea del Sur, 1990-2023
Fuente: Banco Mundial (BN.CAB.XOKA.GD.ZS)
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

wb = Workbook()

# ==================== HOJA 1: DATOS ====================
ws = wb.active
ws.title = "Datos"

# Datos del Banco Mundial (BN.CAB.XOKA.GD.ZS) — descargados febrero 2026
argentina = {
    1990: 3.22, 1991: -0.34, 1992: -2.42, 1993: -3.47, 1994: -4.26,
    1995: -1.98, 1996: -2.49, 1997: -4.14, 1998: -4.84, 1999: -4.21,
    2000: -3.16, 2001: -1.41, 2002: 8.97, 2003: 6.38, 2004: 1.95,
    2005: 2.65, 2006: 2.79, 2007: 2.10, 2008: 1.50, 2009: 2.18,
    2010: -0.38, 2011: -1.01, 2012: -0.39, 2013: -2.38, 2014: -1.74,
    2015: -2.96, 2016: -2.71, 2017: -4.84, 2018: -5.16, 2019: -0.78,
    2020: 0.70, 2021: 1.36, 2022: -0.63, 2023: -3.20
}

corea = {
    1990: -0.96, 1991: -2.36, 1992: -0.79, 1993: 0.42, 1994: -1.00,
    1995: -1.74, 1996: -3.88, 1997: -1.83, 1998: 10.10, 1999: 4.22,
    2000: 1.70, 2001: 0.38, 2002: 0.63, 2003: 1.55, 2004: 3.56,
    2005: 1.26, 2006: 0.19, 2007: 0.86, 2008: 0.16, 2009: 3.37,
    2010: 2.34, 2011: 1.27, 2012: 3.65, 2013: 5.39, 2014: 5.34,
    2015: 6.83, 2016: 6.20, 2017: 4.40, 2018: 4.25, 2019: 3.41,
    2020: 4.35, 2021: 4.39, 2022: 1.44, 2023: 1.78
}

anios = list(range(1990, 2024))

# Encabezados
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

headers = ["Año", "Argentina (% PIB)", "Corea del Sur (% PIB)"]
for j, h in enumerate(headers):
    cell = ws.cell(row=1, column=j+1, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = border

for i, anio in enumerate(anios):
    ws.cell(row=i+2, column=1, value=anio).border = border

    val_ar = argentina[anio]
    cell_ar = ws.cell(row=i+2, column=2, value=round(val_ar, 2))
    cell_ar.number_format = '0.00'
    cell_ar.border = border
    if val_ar < -3:
        cell_ar.fill = PatternFill(start_color="FECACA", end_color="FECACA", fill_type="solid")
    elif val_ar > 0:
        cell_ar.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

    val_kr = corea[anio]
    cell_kr = ws.cell(row=i+2, column=3, value=round(val_kr, 2))
    cell_kr.number_format = '0.00'
    cell_kr.border = border
    if val_kr < -3:
        cell_kr.fill = PatternFill(start_color="FECACA", end_color="FECACA", fill_type="solid")
    elif val_kr > 0:
        cell_kr.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

ws.cell(row=37, column=1, value="Verde = superávit (CC > 0)   |   Rojo = déficit en zona de riesgo (CC < -3%)")
ws.cell(row=37, column=1).font = Font(italic=True, color="666666")
ws.cell(row=38, column=1, value="Fuente: Banco Mundial — Current account balance (% of GDP) [BN.CAB.XOKA.GD.ZS]")
ws.cell(row=38, column=1).font = Font(italic=True, color="666666")

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 22
ws.column_dimensions['C'].width = 24


# ==================== HOJA 2: CÁLCULOS ====================
ws2 = wb.create_sheet("Cálculos")

ws2.cell(row=1, column=1, value="Conteo de años con superávit (CC > 0) — Período 2000-2023").font = Font(bold=True, size=13)

# Tabla de conteo
for j, h in enumerate(["País", "Años con superávit", "Total años", "% con superávit"]):
    cell = ws2.cell(row=3, column=j+1, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = border

# Argentina 2000-2023
ar_superavit = sum(1 for a in range(2000, 2024) if argentina[a] > 0)
kr_superavit = sum(1 for a in range(2000, 2024) if corea[a] > 0)

for i, (pais, sup) in enumerate([("Argentina", ar_superavit), ("Corea del Sur", kr_superavit)]):
    row = 4 + i
    ws2.cell(row=row, column=1, value=pais).border = border
    ws2.cell(row=row, column=1).font = Font(bold=True)
    ws2.cell(row=row, column=2, value=sup).border = border
    ws2.cell(row=row, column=2).font = Font(bold=True, size=14)
    ws2.cell(row=row, column=3, value=24).border = border
    pct = round(sup / 24 * 100, 0)
    ws2.cell(row=row, column=4, value=f"{int(pct)}%").border = border

# Detalle Argentina
ws2.cell(row=7, column=1, value="Detalle Argentina — años con superávit:").font = Font(bold=True, size=11)
ar_sup_years = [a for a in range(2000, 2024) if argentina[a] > 0]
ws2.cell(row=8, column=1, value=", ".join(str(a) for a in ar_sup_years))

ws2.cell(row=10, column=1, value="Detalle Argentina — años en zona de riesgo (CC < -3%):").font = Font(bold=True, size=11, color="CC0000")
ar_risk_years = [(a, argentina[a]) for a in range(2000, 2024) if argentina[a] < -3]
for i, (anio, val) in enumerate(ar_risk_years):
    ws2.cell(row=11+i, column=1, value=anio)
    ws2.cell(row=11+i, column=2, value=round(val, 2))
    ws2.cell(row=11+i, column=2).font = Font(color="CC0000")

row_resp = 11 + len(ar_risk_years) + 2
ws2.cell(row=row_resp, column=1, value="RESPUESTA CORRECTA: b)").font = Font(bold=True, size=12, color="1F4E79")
ws2.cell(row=row_resp+1, column=1, value="Corea: 24/24 años con superávit (aprendió de 1997, nunca más déficit).")
ws2.cell(row=row_resp+2, column=1, value="Argentina: 10/24 años (superávit 2002-2009 + 2020-21, después vuelve al ciclo de déficit).")
ws2.cell(row=row_resp+1, column=1).font = Font(italic=True)
ws2.cell(row=row_resp+2, column=1).font = Font(italic=True)

ws2.column_dimensions['A'].width = 52
ws2.column_dimensions['B'].width = 22
ws2.column_dimensions['C'].width = 14
ws2.column_dimensions['D'].width = 16


# ==================== HOJA 3: GRÁFICO ====================
ws3 = wb.create_sheet("Gráfico")

# Datos para gráfico (solo 2000-2023)
ws3.cell(row=1, column=1, value="Año")
ws3.cell(row=1, column=2, value="Argentina")
ws3.cell(row=1, column=3, value="Corea del Sur")
ws3.cell(row=1, column=4, value="Zona de riesgo (-3%)")

anios_grafico = list(range(2000, 2024))
for i, anio in enumerate(anios_grafico):
    ws3.cell(row=i+2, column=1, value=anio)
    ws3.cell(row=i+2, column=2, value=round(argentina[anio], 2))
    ws3.cell(row=i+2, column=3, value=round(corea[anio], 2))
    ws3.cell(row=i+2, column=4, value=-3)  # línea de referencia

# Gráfico de líneas
chart = LineChart()
chart.title = "Cuenta corriente (% PIB) — Argentina vs Corea del Sur (2000-2023)"
chart.y_axis.title = "% del PIB"
chart.x_axis.title = "Año"
chart.style = 10
chart.width = 30
chart.height = 18

# Agregar las 3 series
for col in range(2, 5):
    data = Reference(ws3, min_col=col, min_row=1, max_row=25)
    chart.add_data(data, titles_from_data=True)

cats = Reference(ws3, min_col=1, min_row=2, max_row=25)
chart.set_categories(cats)

# Colores
colors = ["0EA5E9", "E74C3C", "999999"]
widths = [25000, 25000, 15000]
dashes = [None, None, "dash"]

for idx, series in enumerate(chart.series):
    series.graphicalProperties.line.width = widths[idx]
    series.graphicalProperties.line.solidFill = colors[idx]
    if dashes[idx]:
        series.graphicalProperties.line.dashStyle = dashes[idx]

# Línea de cero
chart.y_axis.crossesAt = 0

ws3.add_chart(chart, "A28")

# Guardar
output = Path(__file__).parent / "U1_P5_cuenta_corriente.xlsx"
wb.save(output)
print(f"Generado: {output}")
print(f"Hojas: Datos, Cálculos, Gráfico")
print(f"\nArgentina superávit 2000-2023: {ar_superavit} de 24 años")
print(f"Corea superávit 2000-2023: {kr_superavit} de 24 años")
print(f"Argentina zona riesgo (<-3%): {[f'{a}: {argentina[a]}%' for a in range(2000,2024) if argentina[a]<-3]}")
