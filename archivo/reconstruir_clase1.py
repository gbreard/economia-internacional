# -*- coding: utf-8 -*-
"""
Reconstruccion de Clase 1 - Economia Internacional
Nueva estructura: 20 slides con narrativa coherente
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Colores del template
AZUL_OSCURO = RGBColor(0x1F, 0x4E, 0x79)
AZUL_CLARO = RGBColor(0x0E, 0xA5, 0xE9)
GRIS_FONDO = RGBColor(0xF3, 0xF4, 0xF6)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
NEGRO = RGBColor(0x11, 0x18, 0x27)
GRIS_TEXTO = RGBColor(0x4B, 0x55, 0x63)
GRIS_CLARO = RGBColor(0xE5, 0xE7, 0xEB)

# Dimensiones
ANCHO = Inches(13.33)
ALTO = Inches(7.5)


def crear_slide_portada(prs, num_clase, titulo, subtitulo, pregunta_guia, total_slides):
    """Crea slide de portada"""
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Fondo gris
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, ANCHO, ALTO)
    background.fill.solid()
    background.fill.fore_color.rgb = GRIS_FONDO
    background.line.fill.background()

    # Titulo principal
    titulo_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.33), Inches(1.0))
    tf = titulo_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Economia Internacional"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = NEGRO
    p.alignment = PP_ALIGN.LEFT

    # Subtitulo
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.1), Inches(12.33), Inches(0.6))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Clase {} - {}".format(num_clase, titulo)
    p.font.size = Pt(20)
    p.font.color.rgb = GRIS_TEXTO

    # Linea de acento
    linea = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(3.8), Inches(2.0), Inches(0.05))
    linea.fill.solid()
    linea.fill.fore_color.rgb = AZUL_CLARO
    linea.line.fill.background()

    # Pregunta guia
    if pregunta_guia:
        guia_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.33), Inches(1.0))
        tf = guia_box.text_frame
        p = tf.paragraphs[0]
        p.text = pregunta_guia
        p.font.size = Pt(14)
        p.font.italic = True
        p.font.color.rgb = GRIS_TEXTO

    # Numeracion
    num_box = slide.shapes.add_textbox(Inches(12.5), Inches(7.0), Inches(0.7), Inches(0.3))
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    p.text = "1/{}".format(total_slides)
    p.font.size = Pt(10)
    p.font.color.rgb = GRIS_TEXTO
    p.alignment = PP_ALIGN.RIGHT

    return slide


def crear_base_slide(prs, num_slide, total_slides, seccion):
    """Crea la estructura base de un slide de contenido"""
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Fondo gris
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, ANCHO, ALTO)
    background.fill.solid()
    background.fill.fore_color.rgb = GRIS_FONDO
    background.line.fill.background()

    # Header azul
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, ANCHO, Inches(0.7))
    header.fill.solid()
    header.fill.fore_color.rgb = AZUL_OSCURO
    header.line.fill.background()

    # Texto header - Clase
    header_clase = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(2), Inches(0.4))
    tf = header_clase.text_frame
    p = tf.paragraphs[0]
    p.text = "Clase 1"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLANCO

    # Texto header - Seccion
    header_sec = slide.shapes.add_textbox(Inches(2.5), Inches(0.22), Inches(8), Inches(0.3))
    tf = header_sec.text_frame
    p = tf.paragraphs[0]
    p.text = seccion
    p.font.size = Pt(12)
    p.font.color.rgb = GRIS_CLARO

    # Area de contenido blanca
    content_area = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.3), Inches(0.9), Inches(12.73), Inches(6.2))
    content_area.fill.solid()
    content_area.fill.fore_color.rgb = BLANCO
    content_area.line.fill.background()

    # Numeracion
    num_box = slide.shapes.add_textbox(Inches(12.5), Inches(7.0), Inches(0.7), Inches(0.3))
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    p.text = "{}/{}".format(num_slide, total_slides)
    p.font.size = Pt(10)
    p.font.color.rgb = GRIS_TEXTO
    p.alignment = PP_ALIGN.RIGHT

    return slide


def agregar_titulo(slide, titulo):
    """Agrega titulo al slide"""
    titulo_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(12.33), Inches(0.6))
    tf = titulo_box.text_frame
    p = tf.paragraphs[0]
    p.text = titulo
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NEGRO


def agregar_contenido_bullets(slide, contenido_items, top=1.8, left=0.5, width=12.33, height=4.5):
    """Agrega contenido con bullets"""
    contenido_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = contenido_box.text_frame
    tf.word_wrap = True

    for i, item in enumerate(contenido_items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        if item.startswith("##"):
            # Subtitulo
            p.text = item[2:].strip()
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = NEGRO
            p.space_before = Pt(12)
        elif item.startswith("-"):
            # Bullet normal
            texto = item[1:].strip()
            p.text = "  " + texto
            p.font.size = Pt(16)
            p.font.color.rgb = NEGRO
            p.space_before = Pt(6)
        elif item.startswith("  -"):
            # Sub-bullet
            texto = item[3:].strip()
            p.text = "      " + texto
            p.font.size = Pt(14)
            p.font.color.rgb = GRIS_TEXTO
            p.space_before = Pt(3)
        elif item == "":
            p.text = ""
            p.space_before = Pt(6)
        else:
            # Texto normal
            p.text = item
            p.font.size = Pt(16)
            p.font.color.rgb = NEGRO
            p.space_before = Pt(6)


def agregar_nota(slide, nota):
    """Agrega nota al pie"""
    nota_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.33), Inches(0.4))
    tf = nota_box.text_frame
    p = tf.paragraphs[0]
    p.text = nota
    p.font.size = Pt(12)
    p.font.italic = True
    p.font.color.rgb = GRIS_TEXTO


def crear_slide_contenido(prs, num_slide, total_slides, seccion, titulo, contenido_items, nota=None):
    """Crea slide de contenido estandar"""
    slide = crear_base_slide(prs, num_slide, total_slides, seccion)
    agregar_titulo(slide, titulo)
    if contenido_items:
        agregar_contenido_bullets(slide, contenido_items)
    if nota:
        agregar_nota(slide, nota)
    return slide


def crear_slide_tabla(prs, num_slide, total_slides, seccion, titulo, headers, rows, nota=None):
    """Crea slide con tabla"""
    slide = crear_base_slide(prs, num_slide, total_slides, seccion)
    agregar_titulo(slide, titulo)

    # Crear tabla
    num_cols = len(headers)
    num_rows = len(rows) + 1
    col_width = 11.5 / num_cols
    table_height = min(0.45 * num_rows, 4.0)

    table = slide.shapes.add_table(
        num_rows, num_cols,
        Inches(0.9), Inches(1.9),
        Inches(11.5), Inches(table_height)
    ).table

    # Headers
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = AZUL_OSCURO
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = BLANCO
        p.alignment = PP_ALIGN.CENTER

    # Rows
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(12)
            p.font.color.rgb = NEGRO

    if nota:
        agregar_nota(slide, nota)

    return slide


def crear_slide_grafico(prs, num_slide, total_slides, seccion, titulo, descripcion_grafico, lectura_items, nota=None):
    """Crea slide preparado para grafico (placeholder)"""
    slide = crear_base_slide(prs, num_slide, total_slides, seccion)
    agregar_titulo(slide, titulo)

    # Placeholder para grafico
    grafico_placeholder = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.7), Inches(1.9),
        Inches(7.5), Inches(4.0)
    )
    grafico_placeholder.fill.solid()
    grafico_placeholder.fill.fore_color.rgb = GRIS_CLARO
    grafico_placeholder.line.color.rgb = AZUL_CLARO

    # Texto en placeholder
    placeholder_text = slide.shapes.add_textbox(Inches(1.5), Inches(3.5), Inches(5.5), Inches(0.5))
    tf = placeholder_text.text_frame
    p = tf.paragraphs[0]
    p.text = "[GRAFICO: {}]".format(descripcion_grafico)
    p.font.size = Pt(14)
    p.font.color.rgb = GRIS_TEXTO
    p.alignment = PP_ALIGN.CENTER

    # Lectura/interpretacion
    if lectura_items:
        lectura_box = slide.shapes.add_textbox(Inches(8.5), Inches(1.9), Inches(4.0), Inches(4.0))
        tf = lectura_box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = "Lectura:"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NEGRO

        for item in lectura_items:
            p = tf.add_paragraph()
            p.text = "  " + item
            p.font.size = Pt(12)
            p.font.color.rgb = NEGRO
            p.space_before = Pt(4)

    if nota:
        agregar_nota(slide, nota)

    return slide


def crear_slide_mapa(prs, num_slide, total_slides, seccion, titulo, descripcion_mapa, ideas_clave, fuente):
    """Crea slide preparado para mapa"""
    slide = crear_base_slide(prs, num_slide, total_slides, seccion)
    agregar_titulo(slide, titulo)

    # Placeholder mapa
    mapa_placeholder = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.7), Inches(1.8),
        Inches(7.0), Inches(4.2)
    )
    mapa_placeholder.fill.solid()
    mapa_placeholder.fill.fore_color.rgb = GRIS_CLARO
    mapa_placeholder.line.color.rgb = AZUL_CLARO

    # Texto placeholder
    placeholder_text = slide.shapes.add_textbox(Inches(1.5), Inches(3.3), Inches(5.0), Inches(1.0))
    tf = placeholder_text.text_frame
    p = tf.paragraphs[0]
    p.text = "[MAPA: {}]".format(descripcion_mapa)
    p.font.size = Pt(14)
    p.font.color.rgb = GRIS_TEXTO
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = "Fuente: {}".format(fuente)
    p.font.size = Pt(10)
    p.font.color.rgb = GRIS_TEXTO
    p.alignment = PP_ALIGN.CENTER

    # Ideas clave
    ideas_box = slide.shapes.add_textbox(Inches(8.0), Inches(1.8), Inches(4.5), Inches(4.2))
    tf = ideas_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Ideas clave:"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NEGRO

    for idea in ideas_clave:
        p = tf.add_paragraph()
        p.text = "  " + idea
        p.font.size = Pt(13)
        p.font.color.rgb = NEGRO
        p.space_before = Pt(6)

    return slide


def main():
    """Construir la presentacion completa"""

    # Crear nueva presentacion
    prs = Presentation()
    prs.slide_width = ANCHO
    prs.slide_height = ALTO

    TOTAL = 20

    # ============================================
    # BLOQUE 1: APERTURA (Slides 1-3)
    # ============================================

    # Slide 1: Portada
    crear_slide_portada(
        prs, 1,
        "El comercio mundial: auge, colapso y preguntas abiertas",
        "Marco de lectura: comercio global + hegemonias + datos (1800-1938)",
        "Pregunta guia: Por que el comercio crecio 40x en un siglo y luego se derrumbo?",
        TOTAL
    )

    # Slide 2: Panorama
    crear_slide_grafico(
        prs, 2, TOTAL,
        "Apertura",
        "Un siglo de crecimiento, treinta anos de destruccion",
        "Indice de comercio mundial 1800-1938 (1913=100)",
        [
            "1800: indice = 2.4",
            "1870: indice = 25 (x10)",
            "1913: indice = 100 (x40)",
            "1918: indice = 75 (-25%)",
            "1929: indice = 133 (pico)",
            "1932: indice = 95 (-29%)"
        ],
        "Este grafico es el mapa del curso - vamos a explicar cada movimiento."
    )

    # Slide 3: Marco analitico
    crear_slide_tabla(
        prs, 3, TOTAL,
        "Marco analitico",
        "Tres fuerzas que explican el comercio",
        ["Fuerza", "Pregunta", "Ejemplo"],
        [
            ["Tecnologia", "Cuanto cuesta mover bienes?", "Vapor, ferrocarril, telegrafo"],
            ["Reglas", "Quien puede comerciar y como?", "Aranceles, tratados, monopolios"],
            ["Poder", "Quien garantiza rutas y pagos?", "Hegemonia, marina, moneda"]
        ],
        "Las tres fuerzas operan juntas. Cuando una falla, el sistema se quiebra."
    )

    # ============================================
    # BLOQUE 2: MEDICION (Slides 4-8)
    # ============================================

    # Slide 4: Indice comercio mundial
    crear_slide_grafico(
        prs, 4, TOTAL,
        "Medicion",
        "Medicion I - Indice de comercio mundial",
        "Indice 1913=100, 1800-1938",
        [
            "Tendencia creciente (s.XIX)",
            "Caida 1914-18 (WWI)",
            "Pico 1929",
            "Colapso 1930-32",
            "Recuperacion parcial 1938"
        ],
        "Fuente: Federico-Tena World Trade Historical Database (FTWTHD)"
    )

    # Slide 5: Exportaciones e importaciones
    crear_slide_grafico(
        prs, 5, TOTAL,
        "Medicion",
        "Medicion II - Exportaciones e importaciones mundiales",
        "Millones de US$ corrientes, 1800-1938",
        [
            "Valores nominales",
            "Crecimiento acelerado post-1870",
            "Caidas en guerras/crisis"
        ],
        "Nota: Para comparar periodos largos, preferir indices (corrigen inflacion)."
    )

    # Slide 6: Participacion regional
    crear_slide_grafico(
        prs, 6, TOTAL,
        "Medicion",
        "Medicion III - Participacion regional en exportaciones",
        "% del total mundial (Europa, Americas, Asia), 1800-1938",
        [
            "Europa domina s.XIX (~65%)",
            "Americas: proveedor primario",
            "Asia: menor pero creciente"
        ],
        "Conexion: Europa (y UK) organiza el comercio del siglo XIX."
    )

    # Slide 7: HHI
    crear_slide_grafico(
        prs, 7, TOTAL,
        "Medicion",
        "Medicion IV - Concentracion regional (HHI)",
        "HHI sobre 5 regiones, 1800-1938",
        [
            "HHI alto = concentrado",
            "HHI bajo = mas parejo",
            "Cae en s.XIX (integracion)",
            "Sube post-1914 (fragmentacion)"
        ],
        "HHI = suma de participaciones al cuadrado. Rango: 0.2 a 1."
    )

    # Slide 8: Cobertura + Definiciones
    crear_slide_contenido(
        prs, 8, TOTAL,
        "Medicion",
        "Cobertura estadistica y definiciones operativas",
        [
            "## Cobertura: cuantas unidades reportan?",
            "- 1800: ~30 polities con datos",
            "- 1913: ~100 polities con datos",
            "- Advertencia: cambios en cobertura afectan niveles",
            "",
            "## Definiciones (para usar con los datos)",
            "- Indice: It = 100 x (Comerciot / Comercio1913)",
            "- Participacion regional: sr,t = Xr,t / Xmundo,t",
            "- Concentracion (HHI): HHIt = Suma (sr,t)^2"
        ],
        "Regla practica: cuando la cobertura cambia rapido, usar indices y mirar rupturas."
    )

    # ============================================
    # BLOQUE 3: AUGE 1815-1913 (Slides 9-12)
    # ============================================

    # Slide 9: Contexto pre-1800
    crear_slide_contenido(
        prs, 9, TOTAL,
        "Historia: Auge",
        "Antes del despegue: el mundo pre-industrial (1500-1800)",
        [
            "## Comercio existia, pero era diferente",
            "- Volumen pequeno: metales preciosos, especias, esclavos",
            "- Rutas controladas por monopolios estatales",
            "  - Casa de Contratacion (Espana)",
            "  - Companias de Indias (Holanda, UK)",
            "- Sin industrializacion = sin produccion en masa",
            "",
            "## En 1800 el indice esta en 2.4",
            "- Casi nada comparado con 1913 (=100)",
            "- El despegue requiere otra tecnologia y otras reglas"
        ],
        "No tenemos datos confiables pre-1800. Desde aqui medimos."
    )

    # Slide 10: Despegue 1815-1870
    crear_slide_contenido(
        prs, 10, TOTAL,
        "Historia: Auge",
        "1815-1870: El despegue - que cambio?",
        [
            "## Tecnologia",
            "- Vapor (1820s): viajes mas rapidos y predecibles",
            "- Ferrocarril: conecta puertos con interior productivo",
            "- Telegrafo (1850s): informacion instantanea de precios",
            "",
            "## Reglas",
            "- UK elimina Corn Laws (1846) = apertura unilateral",
            "- Tratados bilaterales con clausula de nacion mas favorecida",
            "",
            "## Poder",
            "- UK: marina dominante, libra esterlina, City de Londres"
        ],
        "Evidencia: el indice pasa de 2.4 (1800) a 25 (1870) = se multiplica por 10."
    )

    # Slide 11: Primera globalizacion (con mapa)
    crear_slide_mapa(
        prs, 11, TOTAL,
        "Historia: Auge",
        "1870-1913: La primera globalizacion",
        "Red imperial britanica",
        [
            "Rutas maritimas controladas por Royal Navy",
            "Cables submarinos = informacion instantanea",
            "Bases navales en puntos estrategicos",
            "Colonias como proveedores y mercados",
            "El indice se cuadruplica: 25 -> 100"
        ],
        "Navy League, 1922"
    )

    # Slide 12: Hegemonia britanica
    crear_slide_tabla(
        prs, 12, TOTAL,
        "Historia: Auge",
        "Que significa 'hegemonia britanica'?",
        ["Dimension", "Rol de UK", "Efecto"],
        [
            ["Productiva", "Taller del mundo", "Manufactura exportable"],
            ["Financiera", "Libra + City", "Credito y seguros globales"],
            ["Logistica", "Marina mercante", "Rutas predecibles"],
            ["Reglas", "Libre comercio", "Acceso a mercados"]
        ],
        "Hegemonia = proveer bienes publicos internacionales que benefician al hegemon y a otros."
    )

    # ============================================
    # BLOQUE 4: COLAPSO 1914-1945 (Slides 13-15)
    # ============================================

    # Slide 13: WWI
    crear_slide_tabla(
        prs, 13, TOTAL,
        "Historia: Colapso",
        "1914-1918: La guerra destruye el sistema",
        ["Fuerza", "Antes de 1914", "Durante/Despues"],
        [
            ["Tecnologia", "Rutas abiertas", "Bloqueos, submarinos"],
            ["Reglas", "Libre comercio", "Controles, racionamiento"],
            ["Poder", "UK hegemonico", "UK debilitado, fragmentacion"]
        ],
        "Evidencia: el indice cae de 100 (1913) a 75 (1918) = -25% en 4 anos."
    )

    # Slide 14: Gran Depresion
    crear_slide_contenido(
        prs, 14, TOTAL,
        "Historia: Colapso",
        "1929-1932: La Gran Depresion termina la globalizacion",
        [
            "## Secuencia",
            "- Crisis bancaria en EUA -> caida de demanda global",
            "- Respuesta: proteccionismo competitivo (Smoot-Hawley 1930)",
            "- Represalias: cada pais sube aranceles",
            "- Bloques monetarios: area libra, area dolar, area marco",
            "",
            "## Resultado",
            "- Cada pais intenta 'exportar desempleo'",
            "- El comercio colapsa mas rapido que la produccion",
            "- El indice cae de 133 (1929) a 95 (1932) = -29% en 3 anos"
        ],
        "Leccion: las reglas importan - el proteccionismo competitivo destruye comercio."
    )

    # Slide 15: Lecciones del colapso
    crear_slide_contenido(
        prs, 15, TOTAL,
        "Historia: Colapso",
        "Tres lecciones del colapso 1914-1932",
        [
            "## 1. La tecnologia no alcanza",
            "- Los barcos seguian existiendo en 1930",
            "- Pero el comercio cayo igual - faltaban reglas y confianza",
            "",
            "## 2. Las reglas pueden destruir comercio",
            "- Aranceles altos + represalias = espiral descendente",
            "- Sin coordinacion, todos pierden",
            "",
            "## 3. Sin hegemon, el sistema se fragmenta",
            "- UK ya no podia liderar",
            "- EUA no queria hacerlo (aun)",
            "- Resultado: bloques rivales, no sistema global"
        ],
        "Pregunta: estamos viendo algo similar hoy (2018-2024)?"
    )

    # ============================================
    # BLOQUE 5: RECONSTRUCCION 1945-HOY (Slides 16-18)
    # ============================================

    # Slide 16: Posguerra
    crear_slide_tabla(
        prs, 16, TOTAL,
        "Historia: Reconstruccion",
        "1945-1973: EUA asume el rol hegemonico",
        ["Dimension", "UK (pre-1914)", "EUA (post-1945)"],
        [
            ["Moneda", "Libra esterlina", "Dolar (Bretton Woods)"],
            ["Instituciones", "Tratados bilaterales", "FMI, Banco Mundial, GATT"],
            ["Seguridad", "Royal Navy", "US Navy + OTAN"],
            ["Apertura", "Unilateral UK", "Multilateral (rondas GATT)"]
        ],
        "El comercio/PIB mundial supera el nivel de 1913 recien en los anos 1970."
    )

    # Slide 17: Segunda globalizacion
    crear_slide_contenido(
        prs, 17, TOTAL,
        "Historia: Reconstruccion",
        "1973-2008: La segunda globalizacion",
        [
            "## Cambios estructurales",
            "- Fin de Bretton Woods (1971) -> tipos de cambio flotantes",
            "- Contenedores, jets, internet -> costos colapsan",
            "- China entra a OMC (2001) -> 'fabrica del mundo'",
            "",
            "## Resultado",
            "- Comercio crece 2x mas rapido que PIB",
            "- Trade/GDP mundial: 25% (1970) -> 60% (2008)",
            "",
            "## Pero...",
            "- La integracion genera ganadores y perdedores",
            "- Centro vs periferia: quien captura el valor?",
            "- Esto lo veremos con Prebisch (Clase 11)"
        ]
    )

    # Slide 18: Slowbalization
    crear_slide_tabla(
        prs, 18, TOTAL,
        "Historia: Reconstruccion",
        "2008-hoy: El fin de la globalizacion?",
        ["Periodo", "Evento", "Impacto"],
        [
            ["2008-09", "Crisis financiera", "Comercio cae 12%"],
            ["2010-16", "Recuperacion lenta", "Comercio = PIB"],
            ["2018-20", "Guerra comercial", "Aranceles EUA-China"],
            ["2020-22", "Pandemia", "Disrupcion cadenas"],
            ["2022-hoy", "Guerra Ucrania", "Nearshoring"]
        ],
        "Pregunta abierta: Estamos en otro colapso tipo 1914-1932 o es un ajuste?"
    )

    # ============================================
    # BLOQUE 6: CIERRE (Slides 19-20)
    # ============================================

    # Slide 19: Sintesis
    crear_slide_contenido(
        prs, 19, TOTAL,
        "Cierre",
        "Lo que nos llevamos de esta clase",
        [
            "## Un grafico para ver el siglo XIX-XX de un vistazo",
            "- Indice 1800-1938: auge, pico, colapso",
            "",
            "## Un marco para analizar cualquier periodo",
            "- Tecnologia + Reglas + Poder",
            "- Aplicable a 1870, 1930, o 2024",
            "",
            "## Una leccion del colapso 1914-1932",
            "- La integracion no es automatica ni irreversible",
            "- Las reglas y la coordinacion importan",
            "",
            "## Una pregunta abierta",
            "- Que determina si un pais gana o pierde con el comercio?",
            "- Las teorias (Smith, Ricardo, H-O, Prebisch) responden esto"
        ]
    )

    # Slide 20: Proxima clase
    crear_slide_contenido(
        prs, 20, TOTAL,
        "Cierre",
        "Proxima clase: Balanza de pagos y tipo de cambio",
        [
            "## Temas Clase 2",
            "- Cuenta corriente y cuenta financiera",
            "- Tipo de cambio nominal y real",
            "- Que significa 'competitividad externa'?",
            "- Enfoques desarrollistas sobre atraso cambiario",
            "",
            "## Lectura obligatoria",
            "- Krugman, Obstfeld & Melitz - Capitulo 1 (Introduccion)",
            "",
            "## Lectura opcional (traer un grafico para comentar)",
            "- WTO World Trade Report 2024 (resumen ejecutivo)",
            "- UNCTAD Global Trade Update"
        ]
    )

    # Guardar
    prs.save('Clase1_Economia_Internacional_v2.pptx')
    print("Presentacion guardada: Clase1_Economia_Internacional_v2.pptx")
    print("Total slides: {}".format(len(prs.slides)))


if __name__ == "__main__":
    main()
