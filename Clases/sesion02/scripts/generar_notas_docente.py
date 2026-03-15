#!/usr/bin/env python3
"""
Genera un documento Word compacto con todas las notas docente de la Clase 2.
"""
import sys, os, re
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from lxml import etree

# --- Parsear contenido.md ---
script_dir = os.path.dirname(os.path.abspath(__file__))
md_path = os.path.join(script_dir, '..', 'contenido.md')

with open(md_path, 'r', encoding='utf-8') as f:
    raw = f.read()

contenido_match = re.search(r'^## Contenido\s*\n(.*?)(?=^## [^#]|\Z)', raw, re.DOTALL | re.MULTILINE)
if not contenido_match:
    print("ERROR: No se encontró la sección ## Contenido")
    sys.exit(1)

contenido = contenido_match.group(1)

# Parsear slides
slides = []
blocks = re.split(r'^#### ', contenido, flags=re.MULTILINE)

for block in blocks[1:]:
    lines = block.strip().split('\n')
    header = lines[0]
    slide_match = re.match(r'Slide\s+(\d+):\s*(.*)', header)
    if not slide_match:
        continue

    num = int(slide_match.group(1))
    body = '\n'.join(lines[1:])

    tipo_match = re.search(r'\*\*Tipo\*\*:\s*(\w+)', body)
    tipo = tipo_match.group(1) if tipo_match else 'texto'

    titulo_match = re.search(r'\*\*Título\*\*:\s*(.*)', body)
    titulo = titulo_match.group(1).strip() if titulo_match else slide_match.group(2).strip()

    sub_match = re.search(r'\*\*Subtítulo\*\*:\s*(.*)', body)
    subtitulo = sub_match.group(1).strip() if sub_match else ''

    notas_match = re.search(r'\*\*Notas docente\*\*:\s*\n(.*?)(?=\n---|\Z)', body, re.DOTALL)
    notas = ''
    if notas_match:
        notas = notas_match.group(1).strip().replace('```', '').strip()

    slides.append({
        'num': num, 'tipo': tipo, 'titulo': titulo,
        'subtitulo': subtitulo, 'notas': notas
    })

# Secciones
section_map = {}
for sm in re.finditer(r'^### Sección:\s*(.*?)\s*$', contenido, re.MULTILINE):
    pos = sm.end()
    ns = re.search(r'#### Slide\s+(\d+):', contenido[pos:])
    if ns:
        section_map[int(ns.group(1))] = sm.group(1)

# --- Helpers ---
def set_shading(paragraph, color_hex):
    shd = etree.SubElement(paragraph.paragraph_format.element.get_or_add_pPr(), qn('w:shd'))
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)

def compact_spacing(paragraph, before=0, after=0):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

def add_run(paragraph, text, bold=False, italic=False, size=10, color=None):
    r = paragraph.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = 'Calibri'
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r

def is_header_line(line):
    """Detecta si una línea es un header dentro de las notas."""
    s = line.strip()
    if not s or len(s) < 4:
        return False
    if re.search(r'\(\d+\s*minutos?\)', s):
        return True
    # ALL CAPS con al menos 2 palabras, no empieza con - ni número
    if s == s.upper() and not s.startswith('-') and not re.match(r'^\d', s) and ' ' in s:
        return True
    return False

def render_notas_compact(doc, notas_text):
    """Renderiza las notas agrupando líneas consecutivas en un solo párrafo."""
    if not notas_text:
        return

    # Separar en bloques por líneas vacías
    raw_blocks = re.split(r'\n\s*\n', notas_text)

    for block in raw_blocks:
        block = block.strip()
        if not block:
            continue

        lines = block.split('\n')
        first = lines[0].strip()

        # Bloque que empieza con header
        if is_header_line(first):
            p = doc.add_paragraph()
            compact_spacing(p, before=6, after=2)
            add_run(p, first, bold=True, size=10, color=(0x1F, 0x4E, 0x79))
            # Resto del bloque como texto continuo
            rest = [l.strip() for l in lines[1:] if l.strip()]
            if rest:
                p = doc.add_paragraph()
                compact_spacing(p, before=0, after=2)
                for i, line in enumerate(rest):
                    if i > 0:
                        add_run(p, '\n', size=10)
                    # Detectar sub-headers inline
                    if line.endswith(':') and len(line) < 50 and not line.startswith('-'):
                        add_run(p, line, bold=True, size=10, color=(0x2E, 0x75, 0xB6))
                    elif line.startswith('PREGUNTA'):
                        add_run(p, line, bold=True, size=10, color=(0x27, 0xAE, 0x60))
                    elif line.startswith('RESPUESTA'):
                        add_run(p, line, size=10, color=(0x27, 0xAE, 0x60))
                    else:
                        add_run(p, line, size=10)
        else:
            # Bloque normal: todo en un párrafo
            p = doc.add_paragraph()
            compact_spacing(p, before=2, after=2)
            for i, line in enumerate(lines):
                line = line.strip()
                if not line:
                    continue
                if i > 0:
                    add_run(p, '\n', size=10)
                if line.endswith(':') and len(line) < 50 and not line.startswith('-'):
                    add_run(p, line, bold=True, size=10, color=(0x2E, 0x75, 0xB6))
                elif line.startswith('PREGUNTA'):
                    add_run(p, line, bold=True, size=10, color=(0x27, 0xAE, 0x60))
                elif line.startswith('RESPUESTA'):
                    add_run(p, line, size=10, color=(0x27, 0xAE, 0x60))
                elif is_header_line(line):
                    add_run(p, line, bold=True, size=10, color=(0x1F, 0x4E, 0x79))
                else:
                    add_run(p, line, size=10)

# --- Generar Word ---
doc = Document()

# Márgenes compactos
for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

# Estilo base compacto
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
pf = style.paragraph_format
pf.space_before = Pt(0)
pf.space_after = Pt(0)
pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

# --- Portada ---
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, 'NOTAS DOCENTE — CLASE 2', bold=True, size=20, color=(0x1F, 0x4E, 0x79))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, 'Economía Internacional — UMET 2026\n', size=12, color=(0x4B, 0x55, 0x63))
add_run(p, 'Del mercantilismo a Ricardo: productividad, costos relativos y ganancias del comercio\n', size=11, color=(0x4B, 0x55, 0x63))
add_run(p, 'Jueves 26/03/2026 — 18:00 a 22:00 — 34 slides', size=11, color=(0x4B, 0x55, 0x63))

doc.add_paragraph()

# --- Índice compacto ---
p = doc.add_paragraph()
add_run(p, 'ÍNDICE', bold=True, size=12, color=(0x1F, 0x4E, 0x79))
compact_spacing(p, after=4)

current_sec = None
idx_lines = []
for s in slides:
    sec = section_map.get(s['num'], None)
    if sec and sec != current_sec:
        current_sec = sec
        idx_lines.append(('sec', sec))
    idx_lines.append(('slide', s))

p = None
for item_type, item in idx_lines:
    if item_type == 'sec':
        p = doc.add_paragraph()
        compact_spacing(p, before=4, after=1)
        add_run(p, f'▸ {item}', bold=True, size=9, color=(0x2E, 0x75, 0xB6))
    else:
        s = item
        if p is None:
            p = doc.add_paragraph()
            compact_spacing(p, before=0, after=0)
        else:
            add_run(p, '\n', size=9)
        add_run(p, f'   S{s["num"]}: {s["titulo"]}', size=9, color=(0x4B, 0x55, 0x63))
        p = None  # Reset para siguiente sección

doc.add_page_break()

# --- Contenido principal ---
current_sec = None
for s in slides:
    # Sección
    sec = section_map.get(s['num'], None)
    if sec and sec != current_sec:
        current_sec = sec
        p = doc.add_paragraph()
        compact_spacing(p, before=14, after=4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_shading(p, '2E75B6')
        add_run(p, f'  {sec.upper()}  ', bold=True, size=12, color=(0xFF, 0xFF, 0xFF))

    # Header del slide: barra azul oscura con número y título
    p = doc.add_paragraph()
    compact_spacing(p, before=10, after=2)
    set_shading(p, '1F4E79')
    label = f' S{s["num"]}  {s["titulo"]}'
    if s['subtitulo']:
        label += f'  —  {s["subtitulo"]}'
    add_run(p, label, bold=True, size=10, color=(0xFF, 0xFF, 0xFF))

    # Notas
    if s['notas']:
        render_notas_compact(doc, s['notas'])
    else:
        p = doc.add_paragraph()
        compact_spacing(p, before=2, after=2)
        add_run(p, '(Sin notas)', italic=True, size=9, color=(0x9C, 0xA3, 0xAF))

# --- Guardar ---
output_path = os.path.join(script_dir, '..', 'Notas_Docente_Clase2.docx')
doc.save(output_path)
print(f"Documento generado: {output_path}")
print(f"Total slides: {len(slides)}")
print(f"Con notas: {sum(1 for s in slides if s['notas'])}")
