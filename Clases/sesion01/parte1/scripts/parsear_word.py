"""
Script para parsear el Word de la clase y extraer contenido estructurado
"""

from pathlib import Path
import json

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    print("Instalando python-docx...")
    import subprocess
    subprocess.check_call(['pip', 'install', 'python-docx'])
    from docx import Document

# Rutas
BASE_DIR = Path(__file__).parent.parent
WORD_FILE = BASE_DIR / "Clase 1.docx"
OUTPUT_JSON = BASE_DIR / "datos" / "clase1_parseado.json"
OUTPUT_TXT = BASE_DIR / "datos" / "clase1_contenido.txt"


def extraer_contenido_word(filepath):
    """Extrae todo el contenido del Word con su estructura"""

    doc = Document(filepath)

    contenido = {
        "archivo": str(filepath.name),
        "parrafos": [],
        "tablas": [],
        "estructura": []
    }

    # Extraer párrafos con su estilo
    for i, para in enumerate(doc.paragraphs):
        texto = para.text.strip()
        if texto:  # Solo párrafos con contenido
            estilo = para.style.name if para.style else "Normal"

            # Detectar si es título/heading
            es_titulo = "Heading" in estilo or "Title" in estilo or "Título" in estilo

            contenido["parrafos"].append({
                "indice": i,
                "texto": texto,
                "estilo": estilo,
                "es_titulo": es_titulo
            })

            contenido["estructura"].append({
                "tipo": "titulo" if es_titulo else "parrafo",
                "estilo": estilo,
                "texto": texto
            })

    # Extraer tablas
    for i, tabla in enumerate(doc.tables):
        filas = []
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            filas.append(celdas)

        contenido["tablas"].append({
            "indice": i,
            "filas": filas
        })

        contenido["estructura"].append({
            "tipo": "tabla",
            "indice": i,
            "filas": filas
        })

    return contenido


def guardar_txt_legible(contenido, filepath):
    """Guarda una versión de texto plano legible"""

    with open(filepath, 'w', encoding='utf-8') as f:
        f.write("=" * 80 + "\n")
        f.write(f"CONTENIDO PARSEADO DE: {contenido['archivo']}\n")
        f.write("=" * 80 + "\n\n")

        for item in contenido["parrafos"]:
            if item["es_titulo"]:
                f.write("\n" + "=" * 60 + "\n")
                f.write(f"[{item['estilo']}] {item['texto']}\n")
                f.write("=" * 60 + "\n")
            else:
                f.write(f"{item['texto']}\n\n")

        if contenido["tablas"]:
            f.write("\n" + "=" * 80 + "\n")
            f.write("TABLAS\n")
            f.write("=" * 80 + "\n")

            for tabla in contenido["tablas"]:
                f.write(f"\n--- Tabla {tabla['indice'] + 1} ---\n")
                for fila in tabla["filas"]:
                    f.write(" | ".join(fila) + "\n")


if __name__ == "__main__":
    print(f"Parseando: {WORD_FILE}")

    if not WORD_FILE.exists():
        print(f"ERROR: No se encuentra el archivo {WORD_FILE}")
        exit(1)

    contenido = extraer_contenido_word(WORD_FILE)

    # Guardar JSON
    OUTPUT_JSON.parent.mkdir(exist_ok=True)
    with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(contenido, f, ensure_ascii=False, indent=2)
    print(f"JSON guardado: {OUTPUT_JSON}")

    # Guardar TXT legible
    guardar_txt_legible(contenido, OUTPUT_TXT)
    print(f"TXT guardado: {OUTPUT_TXT}")

    # Resumen
    print(f"\nResumen:")
    print(f"  - Párrafos: {len(contenido['parrafos'])}")
    print(f"  - Tablas: {len(contenido['tablas'])}")
    print(f"  - Títulos: {sum(1 for p in contenido['parrafos'] if p['es_titulo'])}")
