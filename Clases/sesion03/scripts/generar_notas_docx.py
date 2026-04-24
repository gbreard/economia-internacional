"""
Genera Notas_Docente_Clase3.docx desde contenido.md
Replica el formato del Word existente:
- Secciones: fondo #2E75B6, texto blanco, centrado, bold, 12pt
- Slides: fondo #1F4E79, texto blanco, bold, 10pt, formato " SN  Título  —  Subtítulo"
- Notas: texto normal 10pt, headers en azul #1F4E79 bold
"""

import sys
import io
import re
import os

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_shading(paragraph, color_hex):
    """Set paragraph background/shading color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), color_hex)
    paragraph.paragraph_format.element.get_or_add_pPr().append(shading)


def parse_contenido_md(filepath):
    """Parse contenido.md and extract slides with their notes."""
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    slides = []
    current_slide = None
    in_notas = False
    notas_lines = []
    current_section = None

    for line in lines:
        stripped = line.rstrip()

        # Section headers
        if stripped.startswith('### Sección:'):
            current_section = stripped.split(':', 1)[1].strip()
            continue

        # Slide header
        if stripped.startswith('#### Slide '):
            # Save previous slide
            if current_slide is not None:
                current_slide['notas_docente'] = '\n'.join(notas_lines).strip()
                slides.append(current_slide)

            # Parse slide number and title
            match = re.match(r'#### Slide (\d+):\s*(.*)', stripped)
            if match:
                num = int(match.group(1))
                title = match.group(2).strip()
            else:
                num = len(slides) + 1
                title = stripped.replace('#### Slide ', '').strip()

            current_slide = {
                'num': num,
                'titulo': title,
                'subtitulo': '',
                'tipo': '',
                'section': current_section,
                'section_changed': False,
            }
            in_notas = False
            notas_lines = []
            continue

        if current_slide is None:
            continue

        # Parse fields
        if stripped.startswith('**Tipo**:'):
            current_slide['tipo'] = stripped.split(':', 1)[1].strip()
        elif stripped.startswith('**Título**:'):
            current_slide['titulo_field'] = stripped.split(':', 1)[1].strip()
        elif stripped.startswith('**Subtítulo**:'):
            current_slide['subtitulo'] = stripped.split(':', 1)[1].strip()
        elif stripped.startswith('**Notas docente**:'):
            in_notas = True
            rest = stripped.split(':', 1)[1].strip()
            if rest:
                notas_lines.append(rest)
        elif stripped == '---':
            in_notas = False
        elif stripped.startswith('**Contenido**:') or stripped.startswith('**Gráfico**:') or \
             stripped.startswith('**Imagen**:') or stripped.startswith('**Fuente**:') or \
             stripped.startswith('**Fecha**:'):
            in_notas = False
        elif in_notas:
            # Strip markdown code block markers
            if stripped.startswith('```'):
                continue
            notas_lines.append(line.rstrip())

    # Save last slide
    if current_slide is not None:
        current_slide['notas_docente'] = '\n'.join(notas_lines).strip()
        slides.append(current_slide)

    # Mark section changes
    prev_section = None
    for s in slides:
        if s['section'] != prev_section and s['section'] is not None:
            s['section_changed'] = True
        prev_section = s['section']

    return slides


def generate_docx(slides, output_path):
    """Generate the Word document with notes."""
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # ===== HEADER =====
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    # empty line

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run('NOTAS DOCENTE — CLASE 3')
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('Economía Internacional — UMET 2026\n')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    run = p.add_run('Del modelo neoclásico a Heckscher-Ohlin: dotaciones, distribución y evidencia\n')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    run = p.add_run('Jueves 23 de abril de 2026')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # ===== INDEX =====
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run('ÍNDICE')
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    current_section = None
    for s in slides:
        if s['section_changed'] and s['section']:
            # Section line in index
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(f"▸ {s['section']}")
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            current_section = s['section']

        # Slide line in index
        titulo = s.get('titulo_field', s['titulo'])
        prefix = '   ' if not s['section_changed'] else '\n   '
        line = f"   S{s['num']}: {titulo}"

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # ===== NOTES FOR EACH SLIDE =====
    current_section = None
    for s in slides:
        notas = s.get('notas_docente', '').strip()
        if not notas:
            continue

        # Section banner
        if s['section_changed'] and s['section']:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            set_shading(p, '2E75B6')
            run = p.add_run(f"  {s['section'].upper()}  ")
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            current_section = s['section']

        # Slide header banner
        titulo = s.get('titulo_field', s['titulo'])
        sub = s.get('subtitulo', '')
        header_text = f" S{s['num']}  {titulo}"
        if sub:
            header_text += f"  —  {sub}"

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        set_shading(p, '1F4E79')
        run = p.add_run(header_text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Notes content - split into paragraphs
        lines = notas.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i]

            # Check if line is a header-like line (ALL CAPS or starts with caps keyword)
            is_header = False
            stripped_line = line.strip()

            if stripped_line and (
                re.match(r'^[A-ZÁÉÍÓÚÑÜ][A-ZÁÉÍÓÚÑÜ\s\-:¿?()0-9×]+$', stripped_line) or
                re.match(r'^(EJEMPLO|PREGUNTA|CONEXIÓN|DATO|RECAPITULACIÓN|CONTEXTO|CLAVE|CONFUSIÓN|ADVERTENCIA|BIOGRAFÍA|EN UNA FRASE|¿POR QUÉ|¿QUÉ|LOS TRES|IR PUNTO|LA PARADOJA|LA CLAVE)', stripped_line)
            ):
                is_header = True

            if not stripped_line:
                i += 1
                continue

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6) if is_header else Pt(2)
            p.paragraph_format.space_after = Pt(2)

            if is_header:
                run = p.add_run(stripped_line)
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            else:
                run = p.add_run(stripped_line)
                run.font.size = Pt(10)

            i += 1

    # Save
    doc.save(output_path)
    print(f"Word generado: {output_path}")
    print(f"Total slides con notas: {sum(1 for s in slides if s.get('notas_docente', '').strip())}")


if __name__ == '__main__':
    script_dir = os.path.dirname(os.path.abspath(__file__))
    base_dir = os.path.dirname(script_dir)
    contenido_path = os.path.join(base_dir, 'contenido.md')
    output_path = os.path.join(base_dir, 'Notas_Docente_Clase3.docx')

    slides = parse_contenido_md(contenido_path)
    print(f"Slides parseados: {len(slides)}")
    generate_docx(slides, output_path)
