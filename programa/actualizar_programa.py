"""
Actualiza Programa de la Asignatura.docx con la nueva estructura de 12 sesiones.
Cambios:
- Sección 1: datos generales (horas, frecuencia, duración)
- Sección 4: contenidos mínimos (ciclo del producto, movilidad de factores)
- Sección 5: programa analítico (Vernon en U3, movilidad de factores en U5)
- Sección 6: metodología (sesiones de 4 horas, coloquio)
- Sección 7: evaluación (MC async + coloquio + recuperatorio)
- Sección 8: bibliografía (Vernon 1966 en U3)
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from copy import deepcopy
from pathlib import Path

def insert_para_after(ref_element, text):
    """Inserta un párrafo nuevo después de ref_element (OxmlElement).
    Copia el estilo del párrafo de referencia. Devuelve el nuevo elemento."""
    new_p = OxmlElement('w:p')
    # Copiar propiedades del párrafo
    pPr = ref_element.find(qn('w:pPr'))
    if pPr is not None:
        new_p.append(deepcopy(pPr))
    # Crear run con texto
    new_r = OxmlElement('w:r')
    # Copiar propiedades del run
    ref_runs = ref_element.findall(qn('w:r'))
    if ref_runs:
        rPr = ref_runs[0].find(qn('w:rPr'))
        if rPr is not None:
            new_r.append(deepcopy(rPr))
    new_t = OxmlElement('w:t')
    new_t.text = text
    new_t.set(qn('xml:space'), 'preserve')
    new_r.append(new_t)
    new_p.append(new_r)
    ref_element.addnext(new_p)
    return new_p


path = Path(r"C:\Users\gbrea\OneDrive\Documentos\UMET\Economia Intenacional\programa\Programa de la Asignatura.docx")
doc = Document(str(path))
p = doc.paragraphs

# ========================================
# VERIFICACIÓN PREVIA
# ========================================
assert "Carga horaria total: 40 horas" in p[9].text, f"[9] inesperado: {p[9].text}"
assert "Frecuencia:" in p[11].text, f"[11] inesperado: {p[11].text}"
assert "Duración estimada:" in p[12].text, f"[12] inesperado: {p[12].text}"
assert "Nuevas teorías del comercio" in p[37].text, f"[37] inesperado: {p[37].text}"
assert "Empresas transnacionales" in p[41].text, f"[41] inesperado: {p[41].text}"
assert "Efectos potenciales" in p[83].text, f"[83] inesperado: {p[83].text}"
assert "Unidad 5:" in p[111].text, f"[111] inesperado: {p[111].text}"
assert "La asignatura combinará:" in p[172].text, f"[172] inesperado: {p[172].text}"
assert "Trabajos prácticos grupales" in p[177].text, f"[177] inesperado: {p[177].text}"
assert "El régimen de evaluación sugerido incluye:" in p[181].text, f"[181] inesperado: {p[181].text}"
assert "Las condiciones específicas" in p[189].text, f"[189] inesperado: {p[189].text}"
assert "Krugman, Obstfeld & Melitz — Caps. 6-8" in p[214].text, f"[214] inesperado: {p[214].text}"
assert "Unidad 5:" in p[229].text, f"[229] inesperado: {p[229].text}"
print("✓ Verificación previa OK — todos los párrafos en posición esperada")

# ========================================
# MODIFICACIONES (no cambian índices)
# ========================================

# --- Sección 1: Datos generales ---
p[9].text = "Carga horaria total: 48 horas"
p[11].text = "Frecuencia: 1 sesión semanal de 4 horas"
p[12].text = "Duración estimada: 12 semanas (12 sesiones de 4 horas)"
print("✓ Sección 1 actualizada")

# --- Sección 4: Contenidos mínimos ---
p[37].text = "Nuevas teorías del comercio: rendimientos crecientes, competencia imperfecta, ciclo del producto, comercio intraindustrial y dumping."
p[41].text = "Movilidad internacional de factores. Empresas transnacionales, inversión extranjera directa y comercio intra-firma."
print("✓ Sección 4 actualizada")

# --- Sección 5: Programa analítico ---
# Título de U5: agregar movilidad de factores
p[111].text = "Unidad 5: Movilidad de factores, empresas transnacionales, inversión extranjera directa y cadenas globales de valor"
print("✓ Sección 5: título U5 actualizado")

# --- Sección 6: Metodología ---
p[172].text = "La asignatura se organiza en sesiones semanales de 4 horas que combinan:"
p[177].text = "Coloquio final grupal: preparación y presentación oral de un caso de inserción internacional (sector, cadena productiva o país), integrando aspectos reales (comercio, CGV) y monetarios (tipo de cambio, restricciones externas)."
print("✓ Sección 6 actualizada")

# --- Sección 7: Evaluación ---
p[181].text = "El régimen de evaluación incluye:"
p[182].text = "Evaluaciones por unidad — multiple choice en plataforma (7 evaluaciones):"
p[183].text = "Al cierre de cada unidad temática se habilita en la plataforma una evaluación de 5 preguntas multiple choice: 3 conceptuales y 2 que requieren procesamiento de datos reales. Plazo de entrega: 1 semana."
p[184].text = "\"Parcial 1\" = promedio de las evaluaciones de las Unidades 1 a 3. \"Parcial 2\" = promedio de las evaluaciones de las Unidades 4 a 7."
p[185].text = "Coloquio final grupal:"
p[186].text = "Presentación y defensa de un caso de inserción internacional (sector, cadena productiva o país), aplicando conceptos del curso. Se asignan los casos con al menos 3 semanas de anticipación. Evaluación: aprobado/desaprobado."
p[187].text = "Recuperatorio:"
p[188].text = "Evaluación MC presencial para quienes hayan reprobado alguna evaluación por unidad."
p[189].text = "Promoción directa: todas las evaluaciones por unidad aprobadas con 7 o más, y coloquio aprobado."
print("✓ Sección 7 actualizada")

# --- Sección 8: Bibliografía ---
# Título U5 en bibliografía
p[229].text = "Unidad 5: Movilidad de factores, empresas transnacionales, IED y cadenas globales de valor"
print("✓ Sección 8: título U5 bib actualizado")

# ========================================
# INSERCIONES (de abajo hacia arriba para no alterar índices)
# ========================================

# Guardar referencias a elementos ANTES de insertar
ref_214 = p[214]._element  # Krugman caps 6-8 en bib U3
ref_189 = p[189]._element  # Última línea de Sección 7
ref_111 = p[111]._element  # Título U5
ref_83 = p[83]._element    # Última línea de Dumping en U3

# 1. Bibliografía U3: insertar Vernon después de Krugman caps 6-8 [214]
insert_para_after(ref_214,
    'Vernon, R. (1966). "International Investment and International Trade in the Product Cycle". '
    'The Quarterly Journal of Economics, Vol. 80, N° 2, pp. 190-207. [EN]')
print("✓ Bibliografía Vernon insertada en U3")

# 2. Sección 7: insertar "Las condiciones..." después de [189]
insert_para_after(ref_189,
    "Las condiciones específicas de regularidad, aprobación y promoción directa "
    "se ajustarán al reglamento de la unidad académica.")
print("✓ Cláusula reglamentaria insertada en Sección 7")

# 3. Programa analítico U5: insertar movilidad de factores después del título [111]
el = insert_para_after(ref_111,
    "Movilidad internacional de factores: marco general.")
el = insert_para_after(el,
    "Del teorema de igualación del precio de los factores a la movilidad efectiva de capital y trabajo.")
el = insert_para_after(el,
    "Movilidad del capital: inversión extranjera directa como forma principal.")
el = insert_para_after(el,
    "Movilidad del trabajo: migración internacional, fuga de cerebros y remesas.")
print("✓ Movilidad de factores insertada en U5")

# 4. Programa analítico U3: insertar Vernon después de Dumping [83]
el = insert_para_after(ref_83,
    "Teoría del ciclo del producto (Vernon):")
el = insert_para_after(el,
    "Innovación en países avanzados, estandarización y traslado de la producción a países de menor costo.")
el = insert_para_after(el,
    "Dinámica temporal del patrón de comercio y localización de la producción internacional.")
print("✓ Vernon insertado en programa analítico U3")

# ========================================
# GUARDAR
# ========================================
doc.save(str(path))
print(f"\n✓ Documento guardado: {path}")

# ========================================
# VERIFICACIÓN POSTERIOR
# ========================================
print("\n=== VERIFICACIÓN POSTERIOR ===\n")
doc2 = Document(str(path))
keywords = [
    'carga horaria', 'frecuencia', 'duración estimada',  # S1
    'ciclo del producto', 'movilidad internacional de factores',  # S4
    'vernon', 'movilidad del capital', 'movilidad del trabajo',  # S5 insertados
    'sesiones semanales', 'coloquio final grupal',  # S6
    'evaluaciones por unidad', 'multiple choice', 'parcial 1', 'recuperatorio',
    'promoción directa', 'las condiciones específicas',  # S7
    'product cycle',  # S8 bib
]
for i, para in enumerate(doc2.paragraphs):
    text_lower = para.text.lower()
    if any(kw in text_lower for kw in keywords):
        print(f"  [{i}] {para.text[:120]}")

print(f"\nTotal de párrafos: {len(doc2.paragraphs)} (antes: 263)")
